import openpyxl
from datetime import datetime
from django.db import transaction
from .models import AttendanceRecord, AttendanceImport
from apps.authentication.models import User

EXPECTED_HEADERS = ["national_id", "date", "check_in", "check_out", "status", "notes"]
VALID_STATUSES   = {"present", "absent", "late", "half_day", "on_leave"}
DATE_FORMAT      = "%Y-%m-%d"
TIME_FORMATS     = ["%H:%M", "%H:%M:%S"]


def _parse_time(value):
    if not value:
        return None
    s = str(value).strip()
    for fmt in TIME_FORMATS:
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            continue
    return None


def generate_template_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(EXPECTED_HEADERS)
    ws.append(["12345678", "2025-01-15", "08:00", "17:00", "present", ""])
    ws.append(["87654321", "2025-01-15", "09:30", "17:00", "late",    "Late arrival"])
    ws.append(["11112222", "2025-01-15", "",      "",      "absent",  "Sick"])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    return wb


def process_attendance_import(import_job: AttendanceImport, file_path: str):
    import_job.status = "processing"
    import_job.save(update_fields=["status"])

    errors = []
    records_to_create = []

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        import_job.status = "failed"
        import_job.errors = [{"row": 0, "error": f"Cannot read file: {e}"}]
        import_job.save()
        return

    if not rows:
        import_job.status = "failed"
        import_job.errors = [{"row": 0, "error": "File is empty"}]
        import_job.save()
        return

    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    if headers != EXPECTED_HEADERS:
        import_job.status = "failed"
        import_job.errors = [{
            "row": 1,
            "error": f"Invalid headers. Expected: {EXPECTED_HEADERS}, got: {headers}"
        }]
        import_job.save()
        return

    user_cache = {
        u.national_id: u
        for u in User.objects.filter(is_active=True).exclude(national_id=None)
    }

    seen_keys = set()

    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue

        national_id, date_str, check_in_str, check_out_str, status_str, notes = (
            row + (None,) * (6 - len(row))
        )[:6]

        if not national_id:
            errors.append({"row": row_idx, "error": "national_id is required"})
            continue

        employee = user_cache.get(str(national_id).strip())
        if not employee:
            errors.append({"row": row_idx, "error": f"Employee '{national_id}' not found"})
            continue

        try:
            date = datetime.strptime(str(date_str).strip(), DATE_FORMAT).date()
        except (ValueError, AttributeError):
            errors.append({"row": row_idx, "error": f"Invalid date '{date_str}' (expected YYYY-MM-DD)"})
            continue

        dup_key = (str(employee.id), str(date))
        if dup_key in seen_keys:
            errors.append({"row": row_idx, "error": f"Duplicate entry for employee '{national_id}' on {date}"})
            continue
        seen_keys.add(dup_key)

        check_in  = _parse_time(check_in_str)
        check_out = _parse_time(check_out_str)

        if check_in_str and check_in is None:
            errors.append({"row": row_idx, "error": f"Invalid check_in time '{check_in_str}'"})
            continue
        if check_out_str and check_out is None:
            errors.append({"row": row_idx, "error": f"Invalid check_out time '{check_out_str}'"})
            continue

        status = str(status_str).strip().lower() if status_str else "present"
        if status not in VALID_STATUSES:
            errors.append({"row": row_idx, "error": f"Invalid status '{status_str}'"})
            continue

        records_to_create.append(AttendanceRecord(
            employee=employee,
            date=date,
            check_in=check_in,
            check_out=check_out,
            status=status,
            notes=str(notes or "").strip(),
            source="excel_import",
            created_by=import_job.uploaded_by,
        ))

    if records_to_create:
        with transaction.atomic():
            AttendanceRecord.objects.bulk_create(
                records_to_create,
                update_conflicts=True,
                update_fields=["check_in", "check_out", "status", "notes"],
                unique_fields=["employee", "date"],
            )

    import_job.total_rows   = len(rows) - 1
    import_job.success_rows = len(records_to_create)
    import_job.failed_rows  = len(errors)
    import_job.errors       = errors
    import_job.status       = "completed"
    import_job.save()
